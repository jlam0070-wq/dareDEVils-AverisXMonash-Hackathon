import io
import os
import re
from pathlib import Path
from typing import Any, Dict, List

import fitz
import pdfplumber
from docx import Document
from openpyxl import load_workbook

from .config import MIN_NATIVE_TEXT_CHARS_PER_PAGE, OCR_DPI


def detect_file_type(path: str | os.PathLike[str]) -> str:
    p = Path(path)
    suffix = p.suffix.lower().lstrip('.')
    with open(p, 'rb') as fh:
        sample = fh.read(4096)
    if suffix == 'pdf' or sample.startswith(b'%PDF'):
        return 'pdf'
    if suffix in {'png', 'jpg', 'jpeg', 'tif', 'tiff'} or sample.startswith((b'\x89PNG', b'\xff\xd8\xff', b'II*\x00', b'MM\x00\x2a')):
        return 'image'
    if suffix in {'docx'} or sample.startswith(b'PK\x03\x04') and b'word/document.xml' in sample:
        return 'docx'
    if suffix in {'xlsx'} or sample.startswith(b'PK\x03\x04') and b'xl/workbook.xml' in sample:
        return 'xlsx'
    if suffix == 'txt' or sample and not sample.startswith(b'PK\x03\x04'):
        return 'txt'
    return 'unknown'


def read_text_file(path: str | os.PathLike[str]) -> str:
    text = Path(path).read_text(encoding='utf-8', errors='replace')
    return text.strip()


def read_docx_file(path: str | os.PathLike[str]) -> str:
    doc = Document(path)
    chunks: List[str] = []
    for paragraph in doc.paragraphs:
        text = (paragraph.text or '').strip()
        if text:
            chunks.append(text)
    for table in doc.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            line = ' | '.join(v for v in values if v)
            if line:
                chunks.append(line)
    return '\n'.join(chunks)


def read_xlsx_file(path: str | os.PathLike[str]) -> str:
    wb = load_workbook(path, data_only=True, read_only=True)
    chunks: List[str] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            values = [str(v).strip() if v is not None else '' for v in row]
            line = ' | '.join(v for v in values if v)
            if line:
                chunks.append(line)
    wb.close()
    return '\n'.join(chunks)


def _pdf_page_text(page: Any) -> str:
    text = page.extract_text() or ''
    tables = page.extract_tables() or []
    table_text: List[str] = []
    for table in tables:
        for row in table:
            row_values = [str(v).strip() if v is not None else '' for v in row]
            table_text.append(' | '.join(v for v in row_values if v))
    return '\n'.join(part for part in [text, *table_text] if part).strip()


def read_pdf_file(path: str | os.PathLike[str]) -> Dict[str, Any]:
    pages: List[Dict[str, Any]] = []
    try:
        with pdfplumber.open(path) as pdf:
            for idx, page in enumerate(pdf.pages, start=1):
                text = _pdf_page_text(page)
                pages.append({'page': idx, 'text': text, 'is_scanned': len(text) < MIN_NATIVE_TEXT_CHARS_PER_PAGE})
    except Exception as exc:  # pdfplumber may fail on malformed PDFs
        return {'pages': [], 'error': str(exc)}
    return {'pages': pages, 'error': None}


def render_pdf_page_to_image(path: str | os.PathLike[str], page_number: int, dpi: int = OCR_DPI) -> bytes:
    doc = fitz.open(path)
    try:
        page = doc[page_number - 1]
        pix = page.get_pixmap(matrix=fitz.Matrix(dpi / 72, dpi / 72))
        return pix.tobytes('png')
    finally:
        doc.close()


def read_native_document(path: str | os.PathLike[str]) -> Dict[str, Any]:
    p = Path(path)
    file_type = detect_file_type(p)
    if file_type == 'txt':
        return {'file_type': 'txt', 'pages': [{'page': 1, 'text': read_text_file(p), 'is_scanned': False, 'source': 'txt'}], 'warnings': []}
    if file_type == 'docx':
        return {'file_type': 'docx', 'pages': [{'page': 1, 'text': read_docx_file(p), 'is_scanned': False, 'source': 'docx'}], 'warnings': []}
    if file_type == 'xlsx':
        return {'file_type': 'xlsx', 'pages': [{'page': 1, 'text': read_xlsx_file(p), 'is_scanned': False, 'source': 'xlsx'}], 'warnings': []}
    if file_type == 'pdf':
        result = read_pdf_file(p)
        if result.get('error'):
            return {'file_type': 'pdf', 'pages': [], 'warnings': [f'pdf read failed: {result["error"]}'], 'error': result['error']}
        return {'file_type': 'pdf', 'pages': result['pages'], 'warnings': [], 'error': None}
    if file_type == 'image':
        return {'file_type': 'image', 'pages': [{'page': 1, 'text': '', 'is_scanned': True, 'source': 'image'}], 'warnings': []}
    return {'file_type': 'unknown', 'pages': [], 'warnings': [f'Unsupported file type {p.suffix or "(no extension)"}'], 'error': 'unsupported_file_type'}
